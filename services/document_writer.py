import os
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph
)

from reportlab.lib.styles import (
    ParagraphStyle
)

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont



def save_document(
    text: str,
    original_path: str,
    output_path: str
):

    extension = os.path.splitext(
        original_path
    )[1].lower()


    if extension == ".txt":

        save_txt(text, output_path)


    elif extension == ".docx":

        save_docx(text, output_path)


    elif extension == ".xlsx":

        save_xlsx(text, output_path)


    elif extension == ".pptx":

        save_pptx(text, output_path)


    elif extension == ".pdf":

        save_pdf(text, output_path)


    else:

        raise ValueError(
            "Формат документа не поддерживается."
        )



def save_txt(
    text: str,
    path: str
):

    with open(
        path,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(text)



def save_docx(
    text: str,
    path: str
):

    document = Document()

    for line in text.splitlines():

        document.add_paragraph(
            line
        )

    document.save(path)



def save_xlsx(
    text: str,
    path: str
):

    workbook = Workbook()

    sheet = workbook.active


    for row, line in enumerate(
        text.splitlines(),
        start=1
    ):

        sheet.cell(
            row=row,
            column=1
        ).value = line


    workbook.save(path)



def save_pptx(
    text: str,
    path: str
):

    presentation = Presentation()

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = "Translation"

    slide.placeholders[1].text = text

    presentation.save(path)


def find_font():
    """
    Ищет TTF-шрифт внутри проекта.
    Шрифт fonts/Arial.ttf копируется в Docker вместе с приложением.
    """

    project_root = Path(__file__).resolve().parent.parent

    font_path = project_root / "fonts" / "Arial.ttf"

    if font_path.is_file():
        return str(font_path)

    raise FileNotFoundError(
        f"Не найден шрифт: {font_path}"
    )


def save_pdf(
    text: str,
    path: str
):
    font_path = find_font()

    pdfmetrics.registerFont(
        TTFont(
            "Arial",
            font_path
        )
    )

    document = SimpleDocTemplate(str(path))

    style = ParagraphStyle(
        "Default",
        fontName="Arial",
        fontSize=12,
        leading=16
    )

    content = []

    for line in text.splitlines():

        if line.strip():

            content.append(
                Paragraph(
                    line,
                    style
                )
            )

    document.build(
        content
    )
