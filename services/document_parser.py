from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from openpyxl import load_workbook


SUPPORTED_FORMATS = (
    ".txt",
    ".docx",
    ".pdf",
    ".pptx",
    ".xlsx"
)


def read_document(path: str) -> str:
    """
    Возвращает текст из документа.
    """

    extension = Path(path).suffix.lower().lstrip(".")

    readers = {
        "txt": read_txt,
        "docx": read_docx,
        "pdf": read_pdf,
        "pptx": read_pptx,
        "xlsx": read_xlsx
    }

    reader = readers.get(extension)

    if reader is None:
        raise ValueError("Формат документа не поддерживается.")

    return reader(path)


def read_txt(path: str) -> str:

    with open(path, "r", encoding="utf-8") as file:
        return file.read()


def read_docx(path: str) -> str:

    document = Document(path)

    text = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n".join(text)


def read_pdf(path: str) -> str:

    try:

        reader = PdfReader(path)

        text = []

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text.append(page_text)

        return "\n".join(text)

    except Exception as error:

        raise ValueError(
            f"Не удалось прочитать PDF: {error}"
        )


def read_pptx(path: str) -> str:

    presentation = Presentation(path)

    text = []

    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                value = shape.text.strip()

                if value:
                    text.append(value)

    return "\n".join(text)


def read_xlsx(path: str) -> str:

    workbook = load_workbook(path)

    text = []

    for sheet in workbook.worksheets:

        for row in sheet.iter_rows():

            values = [
                str(cell.value)
                for cell in row
                if cell.value is not None
            ]

            if values:
                text.append(" ".join(values))

    return "\n".join(text)